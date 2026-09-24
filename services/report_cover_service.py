"""توليد صفحة غلاف تقرير أكاديمي — ثلاثة نماذج (مرتب / رسمي / جامعي) + عربي / إنجليزي."""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from services.text_shape import has_arabic, set_paragraph_direction, set_run_font, shape_for_pdf

logger = logging.getLogger(__name__)

TemplateKind = Literal["modern", "formal", "academic"]
LangKind = Literal["ar", "en"]

ALL_TEMPLATES: tuple[TemplateKind, ...] = ("modern", "formal", "academic")

# ── ألوان ──
MODERN_PRIMARY = (0x0F / 255, 0x17 / 255, 0x2A / 255)
MODERN_ACCENT = (0x25 / 255, 0x63 / 255, 0xEB / 255)
FORMAL_PRIMARY = (0x1A / 255, 0x1A / 255, 0x2E / 255)
FORMAL_ACCENT = (0x4A / 255, 0x4A / 255, 0x68 / 255)
MUTED = (0x64 / 255, 0x74 / 255, 0x8B / 255)

GEO_BLACK = (0.07, 0.07, 0.07)
GEO_MUSTARD = (0.91, 0.72, 0.13)
GEO_MUSTARD_DK = (0.72, 0.54, 0.08)

PRO_NAVY = (0.05, 0.16, 0.32)
PRO_RED = (0.86, 0.12, 0.17)
PRO_INK = (0.08, 0.12, 0.22)

RIB_NAVY = (0.10, 0.22, 0.62)
RIB_MID = (0.18, 0.40, 0.86)
RIB_LIGHT = (0.32, 0.58, 0.95)
RIB_INK = (0.12, 0.14, 0.18)

ACAD_NAVY = (0.10, 0.22, 0.48)
ACAD_RED = (0.78, 0.10, 0.12)
ACAD_GREY = (0.88, 0.88, 0.90)
ACAD_INK = (0.08, 0.08, 0.08)


@dataclass(frozen=True)
class CoverStrings:
    subject: str
    prepared_by: str
    prepared_by_plural: str
    supervisor: str
    date: str
    logo_placeholder: str
    badge: str
    titled: str
    ministry: str
    supervision: str


STRINGS: dict[str, CoverStrings] = {
    "ar": CoverStrings(
        subject="المادة",
        prepared_by="إعداد الطالب",
        prepared_by_plural="إعداد الطلبة",
        supervisor="الإشراف الأكاديمي",
        date="التاريخ",
        logo_placeholder="شعار الجامعة",
        badge="تقرير أكاديمي",
        titled="التقرير بعنوان",
        ministry="وزارة التعليم العالي والبحث العلمي",
        supervision="إشراف",
    ),
    "en": CoverStrings(
        subject="Course",
        prepared_by="Submitted by",
        prepared_by_plural="Submitted by",
        supervisor="Academic Supervisor",
        date="Date",
        logo_placeholder="University Logo",
        badge="ACADEMIC REPORT",
        titled="Report titled",
        ministry="Ministry of Higher Education and Scientific Research",
        supervision="Supervised by",
    ),
}


@dataclass
class ReportCoverData:
    full_name: str
    university: str
    department: str
    subject: str
    title: str
    supervisor: str
    date: str
    lang: LangKind = "ar"
    logo_path: Path | None = None

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> ReportCoverData:
        logo = data.get("logo_path")
        lang = data.get("lang", "ar")
        if lang not in ("ar", "en"):
            lang = "ar"
        return cls(
            full_name=_join_student_names(data.get("full_name", ""), lang),
            university=_clean(data.get("university", "")),
            department=_clean(data.get("department", "")),
            subject=_clean(data.get("subject", "")),
            title=_clean(data.get("title", "")),
            supervisor=_clean(data.get("supervisor", "")),
            date=_clean(data.get("date", "")),
            lang=lang,
            logo_path=Path(logo) if logo else None,
        )

    @property
    def labels(self) -> CoverStrings:
        return STRINGS[self.lang]

    @property
    def rtl(self) -> bool:
        return self.lang == "ar"

    def student_names(self) -> list[str]:
        return _split_student_names(self.full_name)

    def names_line(self) -> str:
        names = self.student_names()
        if not names:
            return self.full_name
        sep = "، " if self.rtl else ", "
        return sep.join(names)

    def prepared_label(self) -> str:
        if len(self.student_names()) > 1:
            return self.labels.prepared_by_plural
        return self.labels.prepared_by


def _paren_title(title: str) -> str:
    text = (title or "").strip()
    if not text:
        return "()"
    if (text.startswith("(") and text.endswith(")")) or (
        text.startswith("（") and text.endswith("）")
    ):
        return text
    return f"({text})"


_NAME_SPLIT = re.compile(r"\s*[,،;/|]\s*|\s+-\s+|\s+–\s+|\s+و\s+")


def _split_student_names(text: str) -> list[str]:
    parts = [_clean(part) for part in _NAME_SPLIT.split(text or "")]
    return [part for part in parts if part]


def _join_student_names(text: str, lang: LangKind) -> str:
    names = _split_student_names(text)
    if not names:
        return _clean(text)
    sep = "، " if lang == "ar" else ", "
    return sep.join(names)


def _clean(text: str) -> str:
    text = (text or "").strip()
    text = re.sub(r"\s+", " ", text)
    if text in {
        "❌ إلغاء", "إلغاء", "Cancel", "cancel",
        "🔙 القائمة الرئيسية", "🏠 القائمة الرئيسية",
        "⏭️ بدون شعار", "⏭️ No logo", "بدون شعار",
    }:
        return ""
    return text


def _display(text: str, lang: LangKind) -> str:
    if lang == "ar" and has_arabic(text):
        return shape_for_pdf(text)
    return text


def _cover_font_dirs() -> list[Path]:
    root = Path(__file__).resolve().parent.parent
    return [root / "fonts", root / "data" / "fonts"]


def _register_pdf_fonts(c) -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_dirs = _cover_font_dirs()
    regular_candidates = [
        *(d / "Amiri-Regular.ttf" for d in font_dirs),
        *(d / "NotoNaskhArabic-Regular.ttf" for d in font_dirs),
        *(d / "Cairo-Regular.ttf" for d in font_dirs),
        Path("/usr/share/fonts/truetype/amiri/Amiri-Regular.ttf"),
        Path("/usr/share/fonts/truetype/hosny-amiri/Amiri-Regular.ttf"),
        Path("C:/Windows/Fonts/trado.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf"),
        Path("C:/Windows/Fonts/tahoma.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    bold_candidates = [
        *(d / "Amiri-Bold.ttf" for d in font_dirs),
        *(d / "NotoNaskhArabic-Bold.ttf" for d in font_dirs),
        *(d / "Cairo-Bold.ttf" for d in font_dirs),
        Path("/usr/share/fonts/truetype/amiri/Amiri-Bold.ttf"),
        Path("/usr/share/fonts/truetype/hosny-amiri/Amiri-Bold.ttf"),
        Path("C:/Windows/Fonts/tradbdo.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoNaskhArabic-Bold.ttf"),
        Path("C:/Windows/Fonts/tahomabd.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ]

    def _reg(name: str, paths: list) -> str | None:
        for fp in paths:
            fp = Path(fp)
            if not fp.exists():
                continue
            try:
                if name not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(name, str(fp)))
                return name
            except Exception as e:
                logger.debug("Font %s skip: %s", fp, e)
        return None

    regular = _reg("CoverRegular", regular_candidates) or "Helvetica"
    bold = _reg("CoverBold", bold_candidates) or regular
    c.setFont(regular, 11)
    logger.debug("Cover PDF fonts: regular=%s bold=%s", regular, bold)
    return regular, bold


def _wrap_by_width(c, text: str, font: str, size: int, max_width: float, lang: LangKind) -> list[str]:
    words = text.split()
    if not words:
        return [""]
    lines: list[str] = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        display = _display(trial, lang)
        if c.stringWidth(display, font, size) <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [text]


def _draw_line_center(c, width: float, y: float, text: str, font: str, size: int,
                      color: tuple, lang: LangKind) -> float:
    if not text:
        return y
    display = _display(text, lang)
    c.setFont(font, size)
    c.setFillColorRGB(*color)
    c.drawCentredString(width / 2, y, display)
    return y - size - 10


def _draw_wrapped_center(c, width: float, y: float, text: str, font: str, size: int,
                         color: tuple, lang: LangKind, margin: float = 55) -> float:
    max_w = width - margin * 2
    for line in _wrap_by_width(c, text, font, size, max_w, lang):
        display = _display(line, lang)
        c.setFont(font, size)
        c.setFillColorRGB(*color)
        c.drawCentredString(width / 2, y, display)
        y -= size + 8
    return y


def _draw_student_names_center(c, width: float, y: float, data: ReportCoverData,
                               font: str, size: int, color: tuple,
                               margin: float = 55) -> float:
    names = data.student_names() or [data.full_name]
    for name in names:
        y = _draw_wrapped_center(c, width, y, name, font, size, color, data.lang, margin=margin)
    return y


def _draw_text(c, x: float, y: float, text: str, font: str, size: int,
               color: tuple, lang: LangKind, align: str = "center") -> None:
    if not text:
        return
    display = _display(text, lang)
    c.setFont(font, size)
    c.setFillColorRGB(*color)
    if align == "right":
        c.drawRightString(x, y, display)
    elif align == "left":
        c.drawString(x, y, display)
    else:
        c.drawCentredString(x, y, display)


def _draw_wrapped_at(c, x: float, y: float, text: str, font: str, size: int,
                     color: tuple, lang: LangKind, max_width: float,
                     align: str = "center") -> float:
    for line in _wrap_by_width(c, text, font, size, max_width, lang):
        _draw_text(c, x, y, line, font, size, color, lang, align)
        y -= size + 8
    return y


def _fill_poly(c, pts: list[tuple[float, float]], color: tuple) -> None:
    p = c.beginPath()
    p.moveTo(pts[0][0], pts[0][1])
    for x, y in pts[1:]:
        p.lineTo(x, y)
    p.close()
    c.setFillColorRGB(*color)
    c.drawPath(p, fill=1, stroke=0)


def _stroke_poly(c, pts: list[tuple[float, float]], color: tuple, lw: float = 1.2) -> None:
    p = c.beginPath()
    p.moveTo(pts[0][0], pts[0][1])
    for x, y in pts[1:]:
        p.lineTo(x, y)
    p.close()
    c.setStrokeColorRGB(*color)
    c.setLineWidth(lw)
    c.drawPath(p, fill=0, stroke=1)


def _draw_logo_pdf(c, width: float, y: float, logo_path: Path | None, placeholder: str,
                   font: str, lang: LangKind, *, max_w: float = 120, max_h: float = 120,
                   color: tuple = MUTED, center_x: float | None = None) -> float:
    from reportlab.lib.utils import ImageReader

    cx = width / 2 if center_x is None else center_x
    if logo_path and logo_path.exists():
        try:
            img = ImageReader(str(logo_path))
            iw, ih = img.getSize()
            scale = min(max_w / iw, max_h / ih, 1.0)
            w, h = iw * scale, ih * scale
            c.drawImage(img, cx - w / 2, y - h, w, h, preserveAspectRatio=True, mask="auto")
            return y - h - 24
        except Exception as e:
            logger.warning("Logo PDF failed: %s", e)
    _draw_text(c, cx, y, f"〔 {placeholder} 〕", font, 9, color, lang)
    return y - 22


def _render_pdf_modern(c, width: float, height: float, data: ReportCoverData,
                       font: str, bold: str) -> None:
    lbl = data.labels
    y = height - 45

    c.setFillColorRGB(*MODERN_ACCENT)
    c.rect(0, height - 18, width, 18, fill=1, stroke=0)
    c.setFillColorRGB(*MODERN_PRIMARY)
    c.rect(0, height - 22, width, 4, fill=1, stroke=0)

    y = _draw_logo_pdf(c, width, y, data.logo_path, lbl.logo_placeholder, font, data.lang)
    y = _draw_line_center(c, width, y, data.university, bold, 22, MODERN_PRIMARY, data.lang)
    y = _draw_line_center(c, width, y, data.department, font, 15, MODERN_ACCENT, data.lang)

    y -= 8
    c.setStrokeColorRGB(*MODERN_ACCENT)
    c.setLineWidth(1.2)
    c.line(width * 0.15, y, width * 0.85, y)
    y -= 28

    subj = f"{lbl.subject}: {data.subject}"
    y = _draw_line_center(c, width, y, subj, font, 13, MUTED, data.lang)
    y -= 20
    y = _draw_wrapped_center(c, width, y, data.title, bold, 26, MODERN_PRIMARY, data.lang)
    y -= 30

    names = data.student_names() or [data.full_name]
    box_h = 40 + max(1, len(names)) * 20
    c.setFillColorRGB(0.97, 0.98, 1.0)
    c.roundRect(width * 0.12, y - box_h, width * 0.76, box_h, 6, fill=1, stroke=0)
    c.setStrokeColorRGB(*MODERN_ACCENT)
    c.setLineWidth(0.6)
    c.roundRect(width * 0.12, y - box_h, width * 0.76, box_h, 6, fill=0, stroke=1)

    inner_y = y - 18
    inner_y = _draw_line_center(
        c, width, inner_y, f"{data.prepared_label()}:", font, 12, MUTED, data.lang,
    )
    for name in names:
        inner_y = _draw_line_center(
            c, width, inner_y, name, font, 13, MODERN_PRIMARY, data.lang,
        )
    _draw_line_center(
        c, width, inner_y,
        f"{lbl.supervisor}: {data.supervisor}", font, 12, FORMAL_ACCENT, data.lang,
    )
    y -= box_h + 20

    c.setStrokeColorRGB(*MUTED)
    c.setLineWidth(0.5)
    c.line(width * 0.3, 70, width * 0.7, 70)
    _draw_line_center(
        c, width, 52, f"{lbl.date}: {data.date}", font, 12, MUTED, data.lang,
    )


def _render_pdf_formal(c, width: float, height: float, data: ReportCoverData,
                       font: str, bold: str) -> None:
    lbl = data.labels
    margin = 42

    c.setStrokeColorRGB(*FORMAL_PRIMARY)
    c.setLineWidth(2)
    c.rect(margin, margin, width - margin * 2, height - margin * 2, fill=0, stroke=1)
    c.setLineWidth(0.8)
    c.rect(margin + 8, margin + 8, width - (margin + 8) * 2, height - (margin + 8) * 2, fill=0, stroke=1)

    y = height - margin - 50
    y = _draw_logo_pdf(c, width, y, data.logo_path, lbl.logo_placeholder, font, data.lang)
    y = _draw_line_center(c, width, y, data.university.upper() if data.lang == "en" else data.university,
                          bold, 20, FORMAL_PRIMARY, data.lang)
    y = _draw_line_center(c, width, y, data.department, font, 14, FORMAL_ACCENT, data.lang)

    y -= 10
    c.setStrokeColorRGB(*FORMAL_PRIMARY)
    c.setLineWidth(1)
    c.line(width * 0.25, y, width * 0.75, y)
    y -= 6
    c.line(width * 0.28, y, width * 0.72, y)
    y -= 26

    y = _draw_line_center(c, width, y, f"{lbl.subject}: {data.subject}", font, 13, MUTED, data.lang)
    y -= 24
    y = _draw_wrapped_center(c, width, y, data.title, bold, 24, FORMAL_PRIMARY, data.lang, margin=70)
    y -= 32

    c.setStrokeColorRGB(*MUTED)
    c.line(width * 0.2, y, width * 0.8, y)
    y -= 28

    y = _draw_line_center(c, width, y, f"{data.prepared_label()}", font, 12, MUTED, data.lang)
    y = _draw_student_names_center(c, width, y, data, bold, 16, FORMAL_PRIMARY, margin=70)
    y -= 8
    y = _draw_line_center(c, width, y, f"{lbl.supervisor}", font, 12, MUTED, data.lang)
    y = _draw_line_center(c, width, y, data.supervisor, font, 14, FORMAL_ACCENT, data.lang)

    _draw_line_center(
        c, width, margin + 28, f"{lbl.date}: {data.date}", font, 12, MUTED, data.lang,
    )


def _render_pdf_royal(c, width: float, height: float, data: ReportCoverData,
                      font: str, bold: str) -> None:
    """هندسي — أسود / أبيض / خردلي بشريط قطري."""
    lbl = data.labels
    c.setFillColorRGB(1, 1, 1)
    c.rect(0, 0, width, height, fill=1, stroke=0)

    _fill_poly(c, [(0, 0), (width, 0), (width, 318), (0, 172)], GEO_BLACK)
    _fill_poly(c, [(0, 248), (width, 88), (width, 168), (0, 328)], GEO_MUSTARD)

    c.setStrokeColorRGB(*GEO_MUSTARD)
    c.setLineWidth(3.2)
    c.line(width - 168, height - 36, width + 8, height - 148)
    c.setLineWidth(1.3)
    c.line(width - 148, height - 22, width + 8, height - 126)

    c.setStrokeColorRGB(1, 1, 1)
    c.setLineWidth(1.6)
    c.line(-8, 96, 188, 42)
    c.setLineWidth(1.0)
    c.line(-8, 74, 168, 24)

    c.setFillColorRGB(1, 1, 1)
    x0, y0 = width - 148, 32
    for row in range(8):
        for col in range(4):
            c.circle(x0 + col * 16, y0 + row * 14, 2.05, fill=1, stroke=0)

    _draw_logo_pdf(
        c, width, height - 40, data.logo_path, lbl.logo_placeholder, font, data.lang,
        max_w=70, max_h=70, color=GEO_BLACK, center_x=62,
    )

    rx, max_w = width - 40, 330
    y = height - 128
    y = _draw_wrapped_at(c, rx, y, data.university, font, 11, GEO_BLACK, data.lang, max_w, "right")
    y -= 8
    y = _draw_wrapped_at(c, rx, y, data.title, bold, 22, GEO_BLACK, data.lang, max_w, "right")
    y -= 4
    y = _draw_wrapped_at(c, rx, y, data.department, font, 11, (0.28, 0.28, 0.28), data.lang, max_w, "right")
    y -= 10
    _draw_text(c, rx, y, f"{lbl.subject}: {data.subject}", font, 10, GEO_MUSTARD_DK, data.lang, "right")

    fx = width - 40
    fy = 128
    _draw_wrapped_at(c, fx, fy, f"{lbl.date}: {data.date}", font, 10, (1, 1, 1), data.lang, 280, "right")
    _draw_wrapped_at(c, fx, fy - 22, f"{data.prepared_label()}: {data.names_line()}", font, 10, (1, 1, 1), data.lang, 280, "right")
    _draw_wrapped_at(c, fx, fy - 44, f"{lbl.supervisor}: {data.supervisor}", font, 10, (1, 1, 1), data.lang, 280, "right")


def _render_pdf_emerald(c, width: float, height: float, data: ReportCoverData,
                        font: str, bold: str) -> None:
    """بروفايل — كحلي / أحمر بأشكال هندسية مائلة."""
    lbl = data.labels
    c.setFillColorRGB(1, 1, 1)
    c.rect(0, 0, width, height, fill=1, stroke=0)

    navy_face = [(292, 548), (width + 4, 638), (width + 4, 392), (258, 302)]
    red_top = [(292, 548), (368, 598), (width + 4, 688), (width + 4, 638)]
    _fill_poly(c, navy_face, PRO_NAVY)
    _fill_poly(c, red_top, PRO_RED)
    _stroke_poly(c, [(x + 14, y + 16) for x, y in navy_face], PRO_RED, 1.3)

    bl = [(0, 36), (218, 118), (176, 198), (0, 118)]
    _fill_poly(c, bl, PRO_NAVY)
    c.setStrokeColorRGB(*PRO_RED)
    c.setLineWidth(1.4)
    c.line(10, 24, 228, 108)
    c.line(10, 24, 10, 44)

    for i in range(11):
        t = i / 10
        col = (PRO_RED[0] * (1 - t) + 0.98 * t, PRO_RED[1] * (1 - t) + 0.86 * t, PRO_RED[2] * (1 - t) + 0.86 * t)
        x = 248 + i * 18
        y = 148 + i * 9
        s = 12
        _fill_poly(c, [(x, y), (x + s, y), (x + s / 2, y + s * 0.88)], col)

    if data.rtl:
        tx, align, max_w = 48 + 300, "right", 300
    else:
        tx, align, max_w = 48, "left", 300
    y = height - 268
    y = _draw_wrapped_at(c, tx, y, data.university, bold, 18, PRO_NAVY, data.lang, max_w, align)
    c.setStrokeColorRGB(*PRO_NAVY)
    c.setLineWidth(1.3)
    if data.rtl:
        c.line(tx - 130, y + 6, tx, y + 6)
    else:
        c.line(tx, y + 6, tx + 130, y + 6)
    y -= 12
    y = _draw_wrapped_at(c, tx, y, data.title, bold, 20, PRO_RED, data.lang, max_w, align)
    y -= 12
    y = _draw_wrapped_at(c, tx, y, data.department, font, 12, PRO_NAVY, data.lang, max_w, align)
    y -= 8
    _draw_text(c, tx, y, f"{lbl.subject}: {data.subject}", font, 11, PRO_RED, data.lang, align)

    info_y = 248
    _draw_text(c, tx, info_y, f"{data.prepared_label()}: {data.names_line()}", font, 10, PRO_INK, data.lang, align)
    _draw_text(c, tx, info_y - 18, f"{lbl.supervisor}: {data.supervisor}", font, 10, PRO_INK, data.lang, align)
    _draw_text(c, tx, info_y - 36, f"{lbl.date}: {data.date}", font, 10, PRO_NAVY, data.lang, align)

    _draw_logo_pdf(
        c, width, 92, data.logo_path, lbl.logo_placeholder, font, data.lang,
        max_w=62, max_h=62, color=PRO_NAVY, center_x=width - 68,
    )


def _render_pdf_ivory(c, width: float, height: float, data: ReportCoverData,
                      font: str, bold: str) -> None:
    """شريطي — شريط أزرق مطوي بدرجات متعددة."""
    lbl = data.labels
    c.setFillColorRGB(1, 1, 1)
    c.rect(0, 0, width, height, fill=1, stroke=0)

    _fill_poly(c, [(0, height), (82, height), (82, 478), (0, 452)], RIB_NAVY)
    _fill_poly(c, [(0, 452), (82, 478), (348, 352), (236, 268)], RIB_MID)
    _fill_poly(c, [(82, 478), (118, 466), (348, 352), (318, 338)], RIB_LIGHT)
    _fill_poly(c, [(236, 268), (348, 352), (width, 208), (width, 0), (292, 0)], RIB_NAVY)
    _fill_poly(c, [(310, 236), (width, 168), (width, 0), (370, 0)], RIB_MID)

    cx, cy, r = 286, 392, 44
    c.setFillColorRGB(*RIB_MID)
    c.circle(cx, cy, r, fill=1, stroke=0)
    c.setFillColorRGB(1, 1, 1)
    c.circle(cx, cy, r - 8, fill=1, stroke=0)
    _draw_text(c, cx, cy - 4, data.subject[:16] or lbl.badge, font, 9, RIB_NAVY, data.lang)
    c.setStrokeColorRGB(*RIB_MID)
    c.setLineWidth(1.4)
    c.line(cx - r - 86, cy, cx - r - 10, cy)
    c.line(cx + r + 10, cy, cx + r + 78, cy)

    _draw_logo_pdf(
        c, width, height - 36, data.logo_path, lbl.logo_placeholder, font, data.lang,
        max_w=58, max_h=58, color=RIB_NAVY, center_x=width - 64,
    )
    _draw_wrapped_at(
        c, width - 36, height - 112, data.university, font, 10, RIB_INK, data.lang, 200, "right",
    )

    y = height - 188
    y = _draw_wrapped_at(c, width - 36, y, data.title, bold, 20, RIB_MID, data.lang, 250, "right")
    y -= 6
    _draw_text(c, width - 36, y, data.department, font, 11, RIB_INK, data.lang, "right")

    _draw_text(c, 108, 210, data.prepared_label(), font, 9, RIB_MID, data.lang, "left")
    _draw_text(c, 108, 190, data.names_line(), bold, 13, RIB_INK, data.lang, "left")
    _draw_text(c, 108, 168, f"{lbl.date}: {data.date}", font, 10, RIB_INK, data.lang, "left")

    _draw_text(c, width - 36, 52, lbl.supervisor, font, 9, (1, 1, 1), data.lang, "right")
    _draw_wrapped_at(
        c, width - 36, 32, data.supervisor, bold, 12, (1, 1, 1), data.lang, 220, "right",
    )


def _render_pdf_academic(c, width: float, height: float, data: ReportCoverData,
                         font: str, bold: str) -> None:
    """جامعي — نفس ترتيب واجهة التقرير: إطار مزخرف، شعار يسار، وزارة/جامعة/قسم."""
    lbl = data.labels
    navy = (0x17 / 255, 0x36 / 255, 0x5D / 255)
    red = (1.0, 0.0, 0.0)
    ink = (0.08, 0.08, 0.08)
    grey = (0xD9 / 255, 0xD9 / 255, 0xD9 / 255)
    title_bg = (0xEA / 255, 0xF1 / 255, 0xDD / 255)

    c.setFillColorRGB(1, 1, 1)
    c.rect(0, 0, width, height, fill=1, stroke=0)

    c.setStrokeColorRGB(*navy)
    for inset, lw in ((16, 0.7), (22, 2.8), (30, 0.7)):
        c.setLineWidth(lw)
        c.rect(inset, inset, width - 2 * inset, height - 2 * inset, fill=0, stroke=1)

    inner = 46
    logo_max = 132
    logo_top = height - inner - 4
    logo_cx = inner + 8 + logo_max / 2
    logo_bottom = _draw_logo_pdf(
        c, width, logo_top, data.logo_path, lbl.logo_placeholder, font, data.lang,
        max_w=logo_max, max_h=logo_max, color=navy, center_x=logo_cx,
    )

    header_w = width - inner * 2 - logo_max - 26
    if data.rtl:
        hx, halign = width - inner - 8, "right"
    else:
        hx, halign = inner + 16 + logo_max, "left"
    y = height - inner - 22
    y = _draw_wrapped_at(c, hx, y, lbl.ministry, bold, 14, ink, data.lang, header_w, halign)
    y -= 4
    y = _draw_wrapped_at(c, hx, y, data.university, bold, 16, ink, data.lang, header_w, halign)
    y -= 2
    y = _draw_wrapped_at(c, hx, y, data.department, bold, 14, ink, data.lang, header_w, halign)

    y = min(y, logo_bottom) - 22

    bar_h = 36
    bar_y = y - bar_h
    c.setFillColorRGB(*grey)
    c.rect(inner, bar_y, width - 2 * inner, bar_h, fill=1, stroke=0)
    _draw_text(
        c, width / 2, bar_y + 12,
        f"{lbl.subject}: {data.subject}", bold, 14, ink, data.lang,
    )

    y = bar_y - 38
    _draw_text(c, width / 2, y, lbl.titled, bold, 15, red, data.lang)
    y -= 34

    title = _paren_title(data.title)
    max_title_w = width - 2 * inner - 48
    for line in _wrap_by_width(c, title, bold, 20, max_title_w, data.lang):
        display = _display(line, data.lang)
        tw = c.stringWidth(display, bold, 20)
        c.setFillColorRGB(*title_bg)
        c.rect(width / 2 - tw / 2 - 10, y - 5, tw + 20, 26, fill=1, stroke=0)
        _draw_text(c, width / 2, y, line, bold, 20, ink, data.lang)
        y -= 30
    y -= 22

    _draw_text(c, width / 2, y, f"{data.prepared_label()}:", bold, 14, ink, data.lang)
    y -= 26
    y = _draw_student_names_center(
        c, width, y, data, bold, 18, ink, margin=inner + 24,
    )

    sup_y = inner + 58
    if y < sup_y + 36:
        sup_y = max(inner + 40, y - 48)
    _draw_text(
        c, width / 2, sup_y,
        f"{lbl.supervision}  {data.supervisor}", bold, 14, red, data.lang,
    )
    _draw_text(c, width / 2, inner + 28, data.date, font, 11, navy, data.lang)


_PDF_RENDERERS = {
    "modern": _render_pdf_modern,
    "formal": _render_pdf_formal,
    "academic": _render_pdf_academic,
}


def generate_cover_pdf(
    data: ReportCoverData,
    output_path: Path,
    *,
    template: TemplateKind = "modern",
) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    output_path.parent.mkdir(parents=True, exist_ok=True)
    width, height = A4
    c = canvas.Canvas(str(output_path), pagesize=A4)
    font, bold = _register_pdf_fonts(c)
    renderer = _PDF_RENDERERS.get(template, _render_pdf_modern)
    renderer(c, width, height, data, font, bold)
    c.save()
    return output_path


def _rgb_docx(rgb: tuple) -> RGBColor:
    return RGBColor(int(rgb[0] * 255), int(rgb[1] * 255), int(rgb[2] * 255))


def _docx_font(rtl: bool) -> str:
    return "Traditional Arabic" if rtl else "Arial"


def _add_para(doc: Document, text: str, *, size: int = 14, bold: bool = False,
              color: RGBColor | None = None, rtl: bool = False, space_after: int = 6,
              space_before: int = 0, align=WD_ALIGN_PARAGRAPH.CENTER):
    para = doc.add_paragraph()
    para.alignment = align
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(space_before)
    set_paragraph_direction(para, rtl)
    run = para.add_run(text)
    set_run_font(run, _docx_font(rtl), size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return para


def _hex_rgb(rgb: tuple) -> str:
    return f"{int(rgb[0]*255):02X}{int(rgb[1]*255):02X}{int(rgb[2]*255):02X}"


def _shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell_border(cell, color: str, sz: str = "8"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _cell_para(cell, text: str, *, size: int = 12, bold: bool = False,
               color: RGBColor | None = None, rtl: bool = False, space_after: int = 4):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    set_paragraph_direction(p, rtl)
    run = p.add_run(text)
    set_run_font(run, _docx_font(rtl), size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return p


def _add_cell_para(cell, text: str, *, size: int = 12, bold: bool = False,
                   color: RGBColor | None = None, rtl: bool = False, space_after: int = 4):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    set_paragraph_direction(p, rtl)
    run = p.add_run(text)
    set_run_font(run, _docx_font(rtl), size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return p


def _add_hrule(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run("─" * 36)
    set_run_font(run, "Arial", 10)
    run.font.color.rgb = _rgb_docx(MUTED)


def _add_logo_docx(doc: Document, logo_path: Path | None, placeholder: str, rtl: bool):
    if logo_path and logo_path.exists():
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(14)
            p.add_run().add_picture(str(logo_path), width=Inches(1.6))
            return
        except Exception as e:
            logger.warning("Logo DOCX failed: %s", e)
    _add_para(doc, f"〔 {placeholder} 〕", size=10, color=_rgb_docx(MUTED), rtl=rtl, space_after=14)


def _add_student_names_docx(doc: Document, data: ReportCoverData, *,
                            size: int = 16, color: RGBColor, rtl: bool,
                            space_after: int = 10, label_size: int = 12,
                            label_color: RGBColor | None = None):
    _add_para(
        doc, f"{data.prepared_label()}:", size=label_size,
        color=label_color or color, rtl=rtl, space_after=2,
    )
    names = data.student_names() or [data.full_name]
    for i, name in enumerate(names):
        after = space_after if i == len(names) - 1 else 3
        _add_para(doc, name, size=size, bold=True, color=color, rtl=rtl, space_after=after)


def _set_page_borders(section, color: str = "17365D"):
    sect_pr = section._sectPr
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "thinThickThinMediumGap")
        el.set(qn("w:sz"), "24")
        el.set(qn("w:space"), "24")
        el.set(qn("w:color"), color)
        pg_borders.append(el)
    sect_pr.append(pg_borders)


def _clear_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tbl_pr.append(borders)


def _set_run_shading(run, hex_color: str):
    r_pr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    r_pr.append(shd)


def _render_docx_modern(doc: Document, data: ReportCoverData) -> None:
    lbl = data.labels
    rtl = data.rtl
    primary = _rgb_docx(MODERN_PRIMARY)
    accent = _rgb_docx(MODERN_ACCENT)
    muted = _rgb_docx(MUTED)

    doc.add_paragraph()
    _add_logo_docx(doc, data.logo_path, lbl.logo_placeholder, rtl)
    _add_para(doc, data.university, size=24, bold=True, color=primary, rtl=rtl, space_after=6)
    _add_para(doc, data.department, size=16, color=accent, rtl=rtl, space_after=12)
    _add_hrule(doc)
    _add_para(doc, f"{lbl.subject}: {data.subject}", size=13, color=muted, rtl=rtl, space_after=16)
    _add_para(doc, data.title, size=28, bold=True, color=primary, rtl=rtl, space_after=24)
    _add_student_names_docx(doc, data, size=14, color=primary, rtl=rtl, space_after=6)
    _add_para(doc, f"{lbl.supervisor}: {data.supervisor}", size=13, color=accent, rtl=rtl, space_after=20)
    _add_hrule(doc)
    _add_para(doc, f"{lbl.date}: {data.date}", size=12, color=muted, rtl=rtl, space_after=0)


def _render_docx_formal(doc: Document, data: ReportCoverData) -> None:
    lbl = data.labels
    rtl = data.rtl
    primary = _rgb_docx(FORMAL_PRIMARY)
    accent = _rgb_docx(FORMAL_ACCENT)
    muted = _rgb_docx(MUTED)

    doc.add_paragraph()
    _add_logo_docx(doc, data.logo_path, lbl.logo_placeholder, rtl)
    uni = data.university.upper() if data.lang == "en" else data.university
    _add_para(doc, uni, size=22, bold=True, color=primary, rtl=rtl, space_after=6)
    _add_para(doc, data.department, size=15, color=accent, rtl=rtl, space_after=14)
    _add_para(doc, "═" * 32, size=11, color=muted, rtl=rtl, space_after=14)
    _add_para(doc, f"{lbl.subject}: {data.subject}", size=13, color=muted, rtl=rtl, space_after=18)
    _add_para(doc, data.title, size=26, bold=True, color=primary, rtl=rtl, space_after=22)
    _add_para(doc, "═" * 32, size=11, color=muted, rtl=rtl, space_after=18)
    _add_student_names_docx(doc, data, size=16, color=primary, rtl=rtl, space_after=10, label_color=muted)
    _add_para(doc, lbl.supervisor, size=12, color=muted, rtl=rtl, space_after=2)
    _add_para(doc, data.supervisor, size=14, color=accent, rtl=rtl, space_after=18)
    _add_para(doc, f"{lbl.date}: {data.date}", size=12, color=muted, rtl=rtl, space_after=0)


def _render_docx_royal(doc: Document, data: ReportCoverData) -> None:
    lbl = data.labels
    rtl = data.rtl
    black = _rgb_docx(GEO_BLACK)
    gold = _rgb_docx(GEO_MUSTARD)
    white = RGBColor(0xFF, 0xFF, 0xFF)
    _add_logo_docx(doc, data.logo_path, lbl.logo_placeholder, rtl)
    _add_para(doc, data.university, size=12, color=black, rtl=rtl, space_after=6)
    _add_para(doc, data.title, size=24, bold=True, color=black, rtl=rtl, space_after=6)
    _add_para(doc, data.department, size=13, color=gold, rtl=rtl, space_after=8)
    _add_para(doc, f"{lbl.subject}: {data.subject}", size=12, color=black, rtl=rtl, space_after=16)
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    _shade_cell(cell, _hex_rgb(GEO_BLACK))
    _cell_para(cell, f"{lbl.date}: {data.date}", size=11, color=white, rtl=rtl)
    _add_cell_para(cell, f"{data.prepared_label()}: {data.names_line()}", size=12, color=gold, rtl=rtl)
    _add_cell_para(cell, f"{lbl.supervisor}: {data.supervisor}", size=12, color=white, rtl=rtl)


def _render_docx_emerald(doc: Document, data: ReportCoverData) -> None:
    lbl = data.labels
    rtl = data.rtl
    navy = _rgb_docx(PRO_NAVY)
    red = _rgb_docx(PRO_RED)
    ink = _rgb_docx(PRO_INK)
    _add_para(doc, data.university, size=20, bold=True, color=navy, rtl=rtl, space_after=4)
    _add_para(doc, "────────", size=10, color=navy, rtl=rtl, space_after=4)
    _add_para(doc, data.title, size=22, bold=True, color=red, rtl=rtl, space_after=8)
    _add_para(doc, data.department, size=13, color=navy, rtl=rtl, space_after=6)
    _add_para(doc, f"{lbl.subject}: {data.subject}", size=12, color=red, rtl=rtl, space_after=16)
    _add_para(doc, f"{data.prepared_label()}: {data.names_line()}", size=12, color=ink, rtl=rtl, space_after=4)
    _add_para(doc, f"{lbl.supervisor}: {data.supervisor}", size=12, color=ink, rtl=rtl, space_after=4)
    _add_para(doc, f"{lbl.date}: {data.date}", size=11, color=navy, rtl=rtl, space_after=12)
    _add_logo_docx(doc, data.logo_path, lbl.logo_placeholder, rtl)


def _render_docx_ivory(doc: Document, data: ReportCoverData) -> None:
    lbl = data.labels
    rtl = data.rtl
    mid = _rgb_docx(RIB_MID)
    navy = _rgb_docx(RIB_NAVY)
    ink = _rgb_docx(RIB_INK)
    _add_logo_docx(doc, data.logo_path, lbl.logo_placeholder, rtl)
    _add_para(doc, data.university, size=12, color=ink, rtl=rtl, space_after=10)
    _add_para(doc, data.title, size=22, bold=True, color=mid, rtl=rtl, space_after=6)
    _add_para(doc, data.department, size=13, color=ink, rtl=rtl, space_after=8)
    _add_para(doc, f"{lbl.subject}: {data.subject}", size=12, color=navy, rtl=rtl, space_after=16)
    _add_para(doc, f"{data.prepared_label()}: {data.names_line()}", size=13, color=ink, rtl=rtl, space_after=4)
    _add_para(doc, f"{lbl.date}: {data.date}", size=11, color=ink, rtl=rtl, space_after=14)
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    _shade_cell(cell, _hex_rgb(RIB_NAVY))
    white = RGBColor(0xFF, 0xFF, 0xFF)
    _cell_para(cell, lbl.supervisor, size=10, color=white, rtl=rtl)
    _add_cell_para(cell, data.supervisor, size=13, bold=True, color=white, rtl=rtl)


def _render_docx_academic(doc: Document, data: ReportCoverData) -> None:
    lbl = data.labels
    rtl = data.rtl
    navy = RGBColor(0x17, 0x36, 0x5D)
    red = RGBColor(0xFF, 0x00, 0x00)
    ink = RGBColor(0x0D, 0x0D, 0x0D)
    header_align = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT

    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    _set_page_borders(section)

    table = doc.add_table(rows=1, cols=2)
    _clear_table_borders(table)
    table.columns[0].width = Inches(2.15)
    table.columns[1].width = Inches(5.35)
    left, right = table.rows[0].cells
    left.width = Inches(2.15)
    right.width = Inches(5.35)

    left.text = ""
    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.paragraph_format.space_after = Pt(0)
    if data.logo_path and data.logo_path.exists():
        try:
            lp.add_run().add_picture(str(data.logo_path), width=Inches(1.7))
        except Exception as e:
            logger.warning("Logo DOCX failed: %s", e)
            run = lp.add_run(f"〔 {lbl.logo_placeholder} 〕")
            set_run_font(run, _docx_font(rtl), 10)
            run.font.color.rgb = navy
    else:
        run = lp.add_run(f"〔 {lbl.logo_placeholder} 〕")
        set_run_font(run, _docx_font(rtl), 10)
        run.font.color.rgb = navy

    right.text = ""
    header_items = (
        (lbl.ministry, 14, ink),
        (data.university, 16, ink),
        (data.department, 14, navy),
    )
    for idx, (text, size, color) in enumerate(header_items):
        p = right.paragraphs[0] if idx == 0 else right.add_paragraph()
        p.alignment = header_align
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)
        set_paragraph_direction(p, rtl)
        for old in list(p.runs):
            old.text = ""
        run = p.add_run(text)
        set_run_font(run, _docx_font(rtl), size)
        run.bold = True
        run.font.color.rgb = color

    doc.add_paragraph()

    subj_table = doc.add_table(rows=1, cols=1)
    _clear_table_borders(subj_table)
    scell = subj_table.rows[0].cells[0]
    _shade_cell(scell, "D9D9D9")
    _cell_para(scell, f"{lbl.subject}: {data.subject}", size=16, bold=True, color=ink, rtl=rtl, space_after=2)

    _add_para(doc, lbl.titled, size=16, bold=True, color=red, rtl=rtl, space_before=16, space_after=10)

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_after = Pt(16)
    set_paragraph_direction(title_para, rtl)
    title_run = title_para.add_run(_paren_title(data.title))
    set_run_font(title_run, _docx_font(rtl), 18)
    title_run.bold = True
    title_run.font.color.rgb = ink
    _set_run_shading(title_run, "EAF1DD")

    _add_student_names_docx(
        doc, data, size=18, color=ink, rtl=rtl, space_after=36, label_size=16,
    )
    _add_para(doc, f"{lbl.supervision}  {data.supervisor}", size=16, bold=True, color=red, rtl=rtl, space_before=28, space_after=10)
    _add_para(doc, data.date, size=12, color=navy, rtl=rtl, space_after=0)

_DOCX_RENDERERS = {
    "modern": _render_docx_modern,
    "formal": _render_docx_formal,
    "academic": _render_docx_academic,
}


def generate_cover_docx(
    data: ReportCoverData,
    output_path: Path,
    *,
    template: TemplateKind = "modern",
) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    renderer = _DOCX_RENDERERS.get(template, _render_docx_modern)
    renderer(doc, data)
    doc.save(str(output_path))
    return output_path


TEMPLATE_NAMES = {
    "ar": {
        "modern": "مرتب",
        "formal": "رسمي",
        "academic": "جامعي",
    },
    "en": {
        "modern": "Modern",
        "formal": "Formal",
        "academic": "Academic",
    },
}
